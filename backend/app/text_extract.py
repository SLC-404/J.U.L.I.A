"""
Saca el texto de un escrito (.docx, .doc o .pdf). Lo usan la ingesta de
escritos del despacho y el RAG para mostrar un escrito completo de ejemplo.
"""

import subprocess
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".doc":
        # .doc viejo de Word: macOS trae `textutil`, que lo convierte sin
        # instalar nada.
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True, text=True, timeout=60,
        )
        return result.stdout
    if ext == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    return ""
