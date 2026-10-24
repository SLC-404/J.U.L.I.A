"""
Paso 6: convierte el texto que redacta el agente en un .docx listo para
editar/imprimir en Word.

El modelo responde en Markdown sencillo (# titulos, **negritas**, listas
con - o 1.), asi que aqui hacemos una conversion minima de ese Markdown a
parrafos de Word. No pretende ser un convertidor completo -- solo lo que
de verdad aparece en un escrito legal.

Los archivos se guardan en:
    storage/clientes/{client_id}_{nombre-cliente}/{work_id}_{nombre-servicio}/archivo.docx
"""

import re
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent.parent
STORAGE_DIR = BASE_DIR / "storage"

FONT_NAME = "Arial"
FONT_SIZE = Pt(12)

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\s*(\d+[.)]|[IVXLC]+\.)\s+(.*)$")
_BOLD_RE = re.compile(r"(\*\*.+?\*\*|__.+?__)")

# Formato del despacho (Lic. Juana): el rubro va arriba a la derecha y los
# titulos de seccion ("H E C H O S", "DERECHO", "PROTESTO LO NECESARIO")
# centrados en negritas; despues de "PROTESTO" todo (fecha y firmas) va
# centrado.
_RUBRO_RE = re.compile(
    r"^\s*(EXPEDIENTE|JUICIO|ACTOR|V\.?\s*S\.?|DEMANDAD|TOCA|CARPETA|DELITO|OFENDID|SENTENCIADO|IMPUTADO|ASUNTO)",
    re.I,
)
_PRESENTE_RE = re.compile(r"P\s*R\s*E\s*S\s*E\s*N\s*T\s*E", re.I)
# "PROTESTO LO NECESARIO" aunque venga con letras separadas ("P R O T E S T O")
_PROTESTO_RE = re.compile(r"^\s*P\s*R\s*O\s*T\s*E\s*S\s*T\s*O", re.I)
# Renglones de la firma de la abogada: van alineados a la izquierda (asi lo
# pidio la Lic. Juana); el lugar/fecha y el nombre del cliente, centrados.
_FIRMA_ABOGADA_RE = re.compile(r"ABOGAD[OA]\s+PATRON|^\s*LIC\.|C[ÉE]D(ULA|\.)\s*PROF", re.I)


def _es_titulo(line: str) -> bool:
    """Linea corta toda en mayusculas, ej. 'H E C H O S' o 'DERECHO'."""
    t = line.strip().rstrip(":.")
    letras = [c for c in t if c.isalpha()]
    return 0 < len(t) <= 40 and len(letras) >= 3 and t == t.upper()


def slugify(text: str) -> str:
    """'Demanda de Desahucio' -> 'demanda-de-desahucio' (sin acentos ni simbolos)."""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text[:60] or "sin-nombre"


def work_dir(client_id: int, client_name: str, work_id: int, service_name: str) -> Path:
    path = (
        STORAGE_DIR
        / "clientes"
        / f"{client_id}_{slugify(client_name)}"
        / f"{work_id}_{slugify(service_name)}"
    )
    path.mkdir(parents=True, exist_ok=True)
    return path


def default_filename(service_name: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"{slugify(service_name)}_{stamp}.docx"


def _add_runs(paragraph, text: str) -> None:
    """Agrega el texto al parrafo respetando **negritas**."""
    for part in _BOLD_RE.split(text):
        if not part:
            continue
        if _BOLD_RE.fullmatch(part):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)


def markdown_to_docx(text: str, dest: Path) -> Path:
    """Escribe `text` como .docx en `dest`. Entiende el formato de los
    escritos del despacho (rubro, titulos en mayusculas, firma) y, por si el
    modelo lo usa, Markdown sencillo (# titulos, **negritas**, listas)."""
    doc = Document()
    _setup_styles(doc)

    antes_de_presente = True
    despues_de_protesto = False

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue

        heading = _HEADING_RE.match(line)
        bullet = _BULLET_RE.match(line)
        numbered = _NUMBERED_RE.match(line)

        if _PROTESTO_RE.match(line):
            despues_de_protesto = True

        if heading:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(heading.group(2).strip("*_ ").upper())
            run.bold = True
        elif antes_de_presente and _RUBRO_RE.match(line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(8)
            p.paragraph_format.space_after = Pt(0)
            _add_runs(p, line.strip())
        elif antes_de_presente and not _PRESENTE_RE.search(line):
            # Bloque "C. JUEZ ... DE TENANGO DEL VALLE ..." a la izquierda.
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            _add_runs(p, line.strip())
        elif despues_de_protesto:
            # "ABOGADO PATRONO / LIC. ... / CED. ..." -> un renglon por parte
            for parte in line.strip().split(" / "):
                parte = parte.strip("* ")
                if not parte:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                if _PROTESTO_RE.match(parte):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(12)
                    run = p.add_run("PROTESTO LO NECESARIO")
                    run.bold = True
                    resto = _PROTESTO_RE.sub("", parte)
                    resto = re.sub(r"^\s*L\s*O\s+N\s*E\s*C\s*E\s*S\s*A\s*R\s*I\s*O\s*[.,]?|^\s*L\s*O\s*N\s*E\s*C\s*E\s*S\s*A\s*R\s*I\s*O\s*[.,]?", "", resto, flags=re.I)
                    if resto.strip(" .,"):
                        p2 = doc.add_paragraph()
                        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p2.add_run(resto.strip(" .,"))
                elif _FIRMA_ABOGADA_RE.search(parte):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if parte.upper().startswith("ABOGAD"):
                        p.paragraph_format.space_before = Pt(36)  # espacio para firmar
                    p.add_run(parte)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if parte.isupper() and len(parte.split()) >= 2 and not any(ch.isdigit() for ch in parte):
                        p.paragraph_format.space_before = Pt(36)  # nombre del cliente: espacio para firmar
                    p.add_run(parte)
        elif _es_titulo(line) and not _PRESENTE_RE.search(line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip().strip("*_"))
            run.bold = True
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, bullet.group(1))
        elif numbered:
            # Se deja la numeracion tal cual la escribio el modelo (1., I., etc.)
            # en vez de la automatica de Word, que suele reiniciarse mal.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, f"{numbered.group(1)} {numbered.group(2)}")
        else:
            p = doc.add_paragraph()
            if _PRESENTE_RE.search(line) and len(line.strip()) < 25:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(line.strip())
                run.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if not antes_de_presente:
                    p.paragraph_format.first_line_indent = Cm(2)
                _add_runs(p, line.strip())

        if _PRESENTE_RE.search(line):
            antes_de_presente = False

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest
